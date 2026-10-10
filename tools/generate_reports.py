from pathlib import Path
import json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[1]
fields=json.loads((root/'tools/report-fields.json').read_text())
for stem,item in fields.items():
    doc=Document()
    doc.core_properties.author='Alexander Baranov'
    doc.core_properties.title=item['title']
    doc.core_properties.comments='Original PCBCalc public report template'
    section=doc.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=Cm(2);section.bottom_margin=Cm(2)
    section.left_margin=Cm(2.2);section.right_margin=Cm(2.2)
    for key in ['Normal','Title','Heading 1']:
        style=doc.styles[key];style.font.name='Arial';style.font.color.rgb=RGBColor(0,0,0)
    # Remove the default document theme's decorative title rule.
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    doc.styles['Normal'].font.size=Pt(11)
    doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    doc.styles['Title'].font.size=Pt(22)
    doc.add_paragraph(item['title'],'Title')
    doc.add_paragraph('Отчёт PCBCalc содержит параметры и результаты расчёта, выполненного в приложении. Значения заполняются при экспорте. Единицы измерения соответствуют полям приложения.')
    doc.add_heading('Параметры и результаты',1)
    for i,name in enumerate(item['bookmarks']):
        p=doc.add_paragraph()
        p.paragraph_format.keep_together=True
        labels = {
            'масса_всех_элементов': 'Масса установленных элементов, кг',
            'Длина': 'Длина платы, м', 'Ширина': 'Ширина платы, м', 'Толщина': 'Толщина платы, м',
            'Плотность': 'Плотность материала, кг/м³', 'Модуль_Юнга': 'Модуль Юнга, ×10¹⁰ Па',
            'коэф_учитывающий_массу_элементов': 'Коэффициент массы элементов',
            'коэф_учит_способ_закрепления': 'Коэффициент закрепления',
            'вспом_коэф_учит_способ_закреп_1': 'Коэффициент закрепления 1',
            'вспом_коэф_учит_способ_закреп_2': 'Коэффициент закрепления 2',
            'вспом_коэф_учит_способ_закреп_3': 'Коэффициент закрепления 3',
            'Возбуждающая_частота': 'Возбуждающая частота, Гц',
            'собственная_частота_колебаний': 'Собственная частота, Гц',
            'собств_поделенная_на_возб': 'Отношение собственной и возбуждающей частот',
            'макс_напряжение_изгиба': 'Расчётное напряжение изгиба, МПа',
            'допустимое_напряжение_изгиба': 'Допустимое напряжение изгиба, МПа',
            'Интен_отказ_сис_учет_разб_кол_эл_на_гр': 'Суммарная интенсивность отказов',
            'Средн_время_наработки_до_перв_отказа': 'Средняя наработка до первого отказа',
        }
        label=labels.get(name,name.replace('_',' '))
        p.add_run(label[:1].upper()+label[1:]+': ').bold=True
        start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(i));start.set(qn('w:name'),name)
        end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(i))
        p._p.append(start);p.add_run('—');p._p.append(end)
    doc.add_paragraph('Результаты предназначены для учебного анализа. Применение к реальному устройству требует проверки исходных данных и расчётной модели.')
    doc.save(root/'app/src/main/assets'/f'{stem}.docx')
